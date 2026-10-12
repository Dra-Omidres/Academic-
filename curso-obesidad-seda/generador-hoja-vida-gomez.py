# -*- coding: utf-8 -*-
"""BORRADOR de hoja de vida de la Dra. Ana María Gómez Medina para el expediente del aval.

La docente no ha enviado su hoja de vida. La Dra. Omidres pidió armarla con lo que se
encontrara en línea (12/10/2026). Fuentes:

  · Cargos y afiliación: las afiliaciones de sus propios artículos en PubMed. La de 2025
    la presenta como jefa de la Unidad de Endocrinología del Hospital Universitario San
    Ignacio; las demás, en la Pontificia Universidad Javeriana.
  · Formación: un resumen de búsqueda (ANHI) dice que hizo Medicina Interna y la
    subespecialidad en Endocrinología en la Javeriana. No se pudo abrir la página para
    comprobarlo, así que va POR CONFIRMAR, igual que los años y el título de médica.
  · Publicaciones: PubMed, consultado el 12/10/2026. 69 resultados con su nombre y
    afiliación Javeriana/San Ignacio; se revisaron 60 y se descartaron 4 que firman otras
    personas («Ana Milena», «Ana María Quintero») o que son de otro campo. Se listan sus
    artículos como primera autora más representativos.

SUPUESTO: que la docente del curso («Dra. Ana María Gómez», endocrinóloga de Colombia,
experta en monitoreo de glucosa) es la Dra. Ana María Gómez Medina. Encaja en todo, pero
hay que confirmarlo. Este documento NO va al expediente hasta que ella lo apruebe.
"""
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.platypus import (BaseDocTemplate, PageTemplate, Frame, Paragraph,
                                Spacer, Table, TableStyle, KeepTogether)

TEAL=colors.HexColor('#0D5A62'); GOLD=colors.HexColor('#B8852B')
MINT=colors.HexColor('#C8E1DF'); PALE=colors.HexColor('#EDF5F4')
WARM=colors.HexColor('#FBF1DC'); INK=colors.HexColor('#1A2325')
MUTED=colors.HexColor('#5F7073'); ROJO=colors.HexColor('#A4161A')

st=lambda **k: ParagraphStyle(**k)
S_nom  = st(name='nom', fontName='Helvetica-Bold', fontSize=16.5, leading=19, textColor=TEAL)
S_esp  = st(name='esp', fontName='Helvetica', fontSize=8.8, leading=12, textColor=GOLD)
S_sec  = st(name='sec', fontName='Helvetica-Bold', fontSize=9.6, leading=12, textColor=colors.white)
S_txt  = st(name='txt', fontName='Helvetica', fontSize=8.5, leading=12.2, textColor=INK, alignment=TA_JUSTIFY)
S_item = st(name='item',fontName='Helvetica', fontSize=8.4, leading=11.6, textColor=INK,
            leftIndent=9, firstLineIndent=-9, spaceAfter=2.2)
S_sub  = st(name='sub', fontName='Helvetica-BoldOblique', fontSize=8.4, leading=11, textColor=TEAL,
            spaceBefore=2, spaceAfter=2)
S_aviso= st(name='aviso', fontName='Helvetica', fontSize=7.8, leading=10.8, textColor=INK)
S_pc   = st(name='pc', fontName='Helvetica-Bold', fontSize=8.4, leading=11.6, textColor=ROJO,
            leftIndent=9, firstLineIndent=-9)

PC='<font color="#A4161A"><b>[POR CONFIRMAR]</b></font>'

def marco(canv, doc):
    canv.saveState()
    canv.setFillColor(TEAL); canv.rect(0, A4[1]-9*mm, A4[0], 9*mm, stroke=0, fill=1)
    canv.setFillColor(GOLD); canv.rect(0, A4[1]-10.2*mm, A4[0], 1.2*mm, stroke=0, fill=1)
    canv.setFillColor(colors.white); canv.setFont('Helvetica-Bold', 7.4)
    canv.drawString(15*mm, A4[1]-6.2*mm, 'DRA. ANA MARÍA GÓMEZ MEDINA')
    canv.setFont('Helvetica', 7.4)
    canv.drawRightString(A4[0]-15*mm, A4[1]-6.2*mm, 'Endocrinología · Tecnología en diabetes')
    # marca de agua de borrador
    canv.setFillColor(colors.Color(.64,.09,.10,alpha=.07)); canv.setFont('Helvetica-Bold', 82)
    canv.translate(A4[0]/2, A4[1]/2); canv.rotate(35)
    canv.drawCentredString(0, 0, 'BORRADOR')
    canv.restoreState(); canv.saveState()
    canv.setStrokeColor(MINT); canv.setLineWidth(.5)
    canv.line(15*mm, 11*mm, A4[0]-15*mm, 11*mm)
    canv.setFillColor(MUTED); canv.setFont('Helvetica', 6.8)
    canv.drawString(15*mm, 7.4*mm, 'Hoja de vida · BORRADOR del 12 de octubre de 2026 · pendiente de aprobación por la docente')
    canv.drawRightString(A4[0]-15*mm, 7.4*mm, 'Página %d' % canv.getPageNumber())
    canv.restoreState()

doc=BaseDocTemplate('SEDA_HojaVida_Gomez_BORRADOR.pdf', pagesize=A4,
    leftMargin=15*mm, rightMargin=15*mm, topMargin=15*mm, bottomMargin=14*mm,
    title='Hoja de vida (borrador) — Dra. Ana María Gómez Medina',
    author='Coordinación académica · Curso virtual SEDA')
doc.addPageTemplates([PageTemplate(id='n',
    frames=[Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id='f')],
    onPage=marco)])
W=doc.width

def seccion(titulo):
    t=Table([[Paragraph(titulo, S_sec)]], colWidths=[W])
    t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),TEAL),
        ('LEFTPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),3.4),
        ('BOTTOMPADDING',(0,0),(-1,-1),3.4)]))
    return [Spacer(1,4*mm), t, Spacer(1,2.4*mm)]

def vinetas(items, estilo=S_item):
    return [Paragraph('•&nbsp; '+i, estilo) for i in items]

def doi(d):
    return 'doi: <link href="https://doi.org/%s" color="#0D5A62">%s</link>' % (d, d)

F=[]

# ---------- aviso para la docente ----------
aviso=Table([[Paragraph(
    '<b>Para la Dra. Gómez.</b> Este borrador lo preparó la coordinación académica del curso '
    'con información pública (sus artículos indexados en PubMed). Le pedimos que lo revise: '
    'corrija lo que haga falta, complete lo marcado como <font color="#A4161A"><b>POR CONFIRMAR</b></font> '
    'y confírmenos que podemos presentarlo en su nombre a la Universidad de Cuenca.', S_aviso)]], colWidths=[W])
aviso.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),WARM),('LINEBEFORE',(0,0),(0,-1),2.2,GOLD),
    ('LEFTPADDING',(0,0),(-1,-1),8),('RIGHTPADDING',(0,0),(-1,-1),8),
    ('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
F += [aviso, Spacer(1,5*mm)]

# ---------- portada ----------
F += [Paragraph('Dra. Ana María Gómez Medina', S_nom), Spacer(1,1.5*mm),
      Paragraph('Médica endocrinóloga · Hospital Universitario San Ignacio · Pontificia Universidad Javeriana · Bogotá, Colombia', S_esp)]

F += seccion('PERFIL PROFESIONAL')
F += [Paragraph(
    'Endocrinóloga dedicada a la tecnología aplicada a la diabetes. Su trabajo clínico y de investigación '
    'se centra en el monitoreo continuo de glucosa, las bombas de insulina y los sistemas de asa cerrada '
    'híbrida, la telemedicina y los ecosistemas digitales para el seguimiento ambulatorio, la hipoglucemia '
    'y la variabilidad glucémica, y el manejo de la diabetes en el paciente hospitalizado. Ha participado '
    'en consensos latinoamericanos sobre el uso del monitoreo continuo de glucosa.', S_txt)]

F += seccion('CARGO ACTUAL')
F += vinetas([
    '<b>Jefa de la Unidad de Endocrinología</b> — Hospital Universitario San Ignacio, Pontificia '
    'Universidad Javeriana, Bogotá, Colombia. <font color="#5F7073">(Según la afiliación de su artículo de 2025.)</font>',
    '<b>Profesora</b> — Departamento de Medicina Interna, Facultad de Medicina, Pontificia Universidad Javeriana ' + PC,
    '<b>Grupo de investigación en Diabetes, Obesidad y Nuevas Tecnologías</b> — Pontificia Universidad Javeriana; rol ' + PC,
])

F += seccion('FORMACIÓN ACADÉMICA')
F += [Paragraph('•&nbsp; Título de médica — universidad y año ' + PC, S_item),
      Paragraph('•&nbsp; Especialidad en Medicina Interna — Pontificia Universidad Javeriana; año ' + PC, S_item),
      Paragraph('•&nbsp; Subespecialidad en Endocrinología — Pontificia Universidad Javeriana; año ' + PC, S_item),
      Paragraph('•&nbsp; Otros posgrados o certificaciones ' + PC, S_item)]

F += seccion('PUBLICACIONES INDEXADAS (PubMed) · selección')
F += [Paragraph('Autora o coautora de más de 50 artículos indexados en PubMed con afiliación a la Pontificia '
                'Universidad Javeriana / Hospital Universitario San Ignacio (búsqueda del 12/10/2026). '
                'Selección como primera autora:', S_txt), Spacer(1,1.5*mm)]
F += vinetas([
    'Gómez-Medina AM, Henao-Carillo DC, Villamil-Castañeda LP, Gómez-Quesada Y, Muñoz-Velandia OM, Yepes CA, et al. Insulin DEgludec/LIraglutide versus multiple daily insulin injections in the transition from hospital to outpatient management assessed by continuous glucose monitoring: the DELI transition trial. <i>Diabetologia</i>. 2025;68(9):1899-1907. ' + doi('10.1007/s00125-025-06446-y'),
    'Gómez AM, Henao-Carrillo D, Muñoz-Velandia O, García-Jaramillo M, Hurtado D, Vargas C, et al. Advanced Hybrid Closed-Loop Systems in Perioperative Care: Safety During Anesthesia. <i>Diabetes Technol Ther</i>. 2026;28(2):142-150. ' + doi('10.1177/15209156251376011'),
    'Gómez AM, Henao-Carrillo D, Muñoz-Velandia O, Mina MA, Bermont PA. Virtual vs. In-Person Training for MinMed™ 780G Initiation: Comparable Effectiveness and Safety in Adults with Type 1 Diabetes. <i>Clin Med Insights Endocrinol Diabetes</i>. 2026;19:11795514261471612. ' + doi('10.1177/11795514261471612'),
    'Gomez AM, Henao DC, Muñoz OM, Lucero OD, Cortes D, Del Castillo Cuervo A, et al. Effectiveness of Using a Digital Ecosystem for Ambulatory Diabetes Care in Patients Diagnosed With Type 2 Diabetes Mellitus. <i>Clin Med Insights Endocrinol Diabetes</i>. 2025;18:11795514251349342. ' + doi('10.1177/11795514251349342'),
    'Gómez AM, Henao Carrillo DC, Ré MA, Faradji RN, Flores Caloca O, de la Garza Hernández NE, et al. Recommendations on the use of the flash continuous glucose monitoring system in hospitalized patients with diabetes in Latin America. <i>Diabetol Metab Syndr</i>. 2024;16(1):128. ' + doi('10.1186/s13098-024-01362-4'),
    'Gómez AM, Henao DC, Muñoz OM, Romero DM, León JDS, Jaramillo PE, et al. Temporary Target Versus Suspended Insulin Infusion in Patients with Type 1 Diabetes Using the MiniMed 780G Advanced Closed-Loop Hybrid System During Aerobic Exercise: A Randomized Crossover Clinical Trial. <i>Diabetes Technol Ther</i>. 2024;26(11):823-828. ' + doi('10.1089/dia.2023.0589'),
    'Gómez Medina AM, Henao Carrillo DC, Silva León JD, Gómez González JA, Muñoz Velandia OM, Conde Brahim L, et al. Results From a Virtual Clinic for the Follow-up of Patients Using the Advanced Hybrid Closed-Loop System. <i>J Diabetes Sci Technol</i>. 2023;19(3):778-786. ' + doi('10.1177/19322968231204376'),
    'Gómez AM, Henao D, Parra D, Kerguelen A, Jaramillo P, Gómez Y, et al. Early and sustained increase in time in range 1 year after initiation of hybrid close loop therapy via telemedicine in type 1 diabetes patients. <i>Acta Diabetol</i>. 2023;60(7):943-949. ' + doi('10.1007/s00592-023-02051-w'),
    'Gómez AM, Henao C DC, Rebolledo M, Jaramillo P PE, Muñoz V OM, Niño G LM, et al. Determination of Time in Range Associated With HbA1c ≤7% in a Prospective Cohort of Patients With Type 1 Diabetes Using Continuous Glucose Monitoring for Three Months. <i>J Diabetes Sci Technol</i>. 2022;18(2):345-350. ' + doi('10.1177/19322968221108424'),
    'Gómez AM, Henao DC, Muñoz OM, Aschner P, Yepes CA, Jojoa R, et al. Glycemic control metrics using flash glucose monitoring and hospital complications in patients with COVID-19. <i>Diabetes Metab Syndr</i>. 2021;15(2):499-503. ' + doi('10.1016/j.dsx.2021.02.008'),
    'Gómez AM, Henao DC, Imitola Madero A, Taboada LB, Cruz V, Robledo Gómez MA, et al. Defining High Glycemic Variability in Type 1 Diabetes: Comparison of Multiple Indexes to Identify Patients at Risk of Hypoglycemia. <i>Diabetes Technol Ther</i>. 2019;21(8):430-439. ' + doi('10.1089/dia.2019.0075'),
    'Gómez AM, Marín Carrillo LF, Arévalo Correa CM, Muñoz Velandia OM, Rondón Sepúlveda MA, Silva Herrera JL, et al. Maternal-Fetal Outcomes in 34 Pregnant Women with Type 1 Diabetes in Sensor-Augmented Insulin Pump Therapy. <i>Diabetes Technol Ther</i>. 2017;19(7):417-422. ' + doi('10.1089/dia.2017.0030'),
    'Gómez AM, Henao Carrillo DC, Muñoz Velandia OM. Devices for continuous monitoring of glucose: update in technology. <i>Med Devices (Auckl)</i>. 2017;10:215-224. ' + doi('10.2147/MDER.S110121'),
    'Gomez AM, Alfonso-Cristancho R, Orozco JJ, Lynch PM, Prieto D, Saunders R, et al. Clinical and economic benefits of integrated pump/CGM technology therapy in patients with type 1 diabetes in Colombia. <i>Endocrinol Nutr</i>. 2016;63(9):466-474. ' + doi('10.1016/j.endonu.2016.05.011'),
])

F += seccion('DOCENCIA EN ESTE CURSO')
F += vinetas([
    '<b>Curso virtual de actualización en obesidad, diabetes, nutrición clínica y salud digital</b> — '
    'Sociedad de Endocrinología y Diabetes del Austro (SEDA), Ecuador. Módulo 7 · Clase 2: '
    '<i>Monitoreo digital y tecnologías en diabetes y obesidad</i>.',
])

F += [Spacer(1,4*mm),
      Paragraph('<font color="#5F7073" size="7.2">Fuentes de este borrador: registros de PubMed consultados el '
                '12/10/2026 (publicaciones y afiliaciones). Formación y cargos docentes por confirmar con la docente.</font>', S_txt)]

F_word=list(F)   # doc.build vacía la lista
doc.build(F)

# ---------- versión Word, editable, con el mismo contenido ----------
# Se recorre lo que ya se armó para el PDF, así ambos documentos no se desalinean.
import re
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def _rgb(h): h=h.lstrip('#'); return RGBColor(int(h[:2],16), int(h[2:4],16), int(h[4:],16))

def _sombra(par, color):
    pPr=par._p.get_or_add_pPr(); shd=OxmlElement('w:shd')
    shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),color)
    pPr.append(shd)

def _runs(par, marca, size, color='1A2325', negrita=False):
    """Convierte el marcado de reportlab (<b>, <i>, <font color>, <link>) en runs de Word."""
    marca=marca.replace('&nbsp;',' ').replace('•  ','')
    estado={'b':negrita,'i':False,'c':[color]}
    for trozo in re.split(r'(<[^>]+>)', marca):
        if not trozo: continue
        t=trozo.lower()
        if t in ('<b>',): estado['b']=True; continue
        if t in ('</b>',): estado['b']=negrita; continue
        if t in ('<i>',): estado['i']=True; continue
        if t in ('</i>',): estado['i']=False; continue
        m=re.match(r'<(?:font|link)[^>]*color="#?([0-9A-Fa-f]{6})"', trozo)
        if m: estado['c'].append(m.group(1)); continue
        if t.startswith('<font') or t.startswith('<link'): estado['c'].append(estado['c'][-1]); continue
        if t in ('</font>','</link>'): estado['c'].pop(); continue
        if trozo.startswith('<'): continue
        r=par.add_run(trozo); r.bold=estado['b']; r.italic=estado['i']
        r.font.size=Pt(size); r.font.color.rgb=_rgb(estado['c'][-1]); r.font.name='Arial'

wd=Document()
for sec in wd.sections:
    sec.top_margin=sec.bottom_margin=Cm(1.6); sec.left_margin=sec.right_margin=Cm(1.8)
    h=sec.header.paragraphs[0]
    _runs(h, '<b>DRA. ANA MARÍA GÓMEZ MEDINA</b>   ·   <font color="#A4161A"><b>BORRADOR</b></font> '
             '— pendiente de aprobación por la docente', 8, '0D5A62')
    f=sec.footer.paragraphs[0]
    _runs(f, 'Hoja de vida · BORRADOR del 12 de octubre de 2026', 7.5, '5F7073')

def _par(espacio_antes=0, espacio_despues=3):
    q=wd.add_paragraph(); q.paragraph_format.space_before=Pt(espacio_antes)
    q.paragraph_format.space_after=Pt(espacio_despues); return q

for fl in F_word:
    if isinstance(fl, Table):
        par_rl=fl._cellvalues[0][0]
        if not isinstance(par_rl, Paragraph): par_rl=par_rl[0]
        if par_rl.style.name=='sec':
            q=_par(9,4); q.paragraph_format.keep_with_next=True; _sombra(q,'0D5A62'); _runs(q, ' '+par_rl.text, 10, 'FFFFFF', True)
        else:
            q=_par(0,8); _sombra(q,'FBF1DC'); _runs(q, par_rl.text, 8.5)
    elif isinstance(fl, Paragraph):
        n=fl.style.name
        if n=='nom': _runs(_par(0,1), fl.text, 18, '0D5A62', True)
        elif n=='esp': _runs(_par(0,4), fl.text, 9.5, 'B8852B')
        elif n=='sub':
            q=_par(4,2); q.paragraph_format.keep_with_next=True; _runs(q, fl.text, 9, '0D5A62', True)
        elif n=='item':
            q=wd.add_paragraph(style='List Bullet'); q.paragraph_format.space_after=Pt(3)
            _runs(q, fl.text, 9)
        else:
            q=_par(0,4); q.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; _runs(q, fl.text, 9)
wd.core_properties.title='Hoja de vida (borrador) — Dra. Ana María Gómez Medina'
wd.core_properties.author='Coordinación académica · Curso virtual SEDA'
wd.save('SEDA_HojaVida_Gomez_BORRADOR.docx')
print('generado')
