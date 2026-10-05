# -*- coding: utf-8 -*-
"""Directorio de ponentes para la gestión documental del aval. Herramienta de trabajo
para la Srta. Natalia (RRSS SEDA): a quién le falta qué y cómo contactarla."""
import openpyxl, unicodedata, re
from difflib import SequenceMatcher
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import (BaseDocTemplate, PageTemplate, Frame, Paragraph,
                                Spacer, Table, TableStyle, KeepTogether)

TEAL=colors.HexColor('#0D5A62'); GOLD=colors.HexColor('#B8852B')
MINT=colors.HexColor('#C8E1DF'); PALE=colors.HexColor('#EDF5F4')
WARM=colors.HexColor('#FBF1DC'); INK=colors.HexColor('#1A2325'); MUTED=colors.HexColor('#5F7073')
OK  =colors.HexColor('#2E6B4F'); FALTA=colors.HexColor('#A3342B')
VERDE=colors.HexColor('#EAF4EE')

S=lambda **k: ParagraphStyle(**k)
st_tit=S(name='t', fontName='Helvetica-Bold', fontSize=18, leading=22, textColor=TEAL)
st_sub=S(name='s', fontName='Helvetica', fontSize=9.5, leading=13, textColor=MUTED)
st_sec=S(name='se',fontName='Helvetica-Bold', fontSize=10.5, leading=13, textColor=colors.white)
st_nom=S(name='n', fontName='Helvetica-Bold', fontSize=9, leading=11, textColor=INK)
st_dat=S(name='d', fontName='Helvetica', fontSize=8, leading=10, textColor=MUTED)
st_fal=S(name='f', fontName='Helvetica-Bold', fontSize=8.2, leading=10.5, textColor=FALTA)
st_ok =S(name='o', fontName='Helvetica-Bold', fontSize=8.2, leading=10.5, textColor=OK)
st_hd =S(name='h', fontName='Helvetica-Bold', fontSize=7.4, leading=9, textColor=MUTED)
st_av =S(name='a', fontName='Helvetica', fontSize=8.8, leading=11.8, textColor=INK)
st_big=S(name='b', fontName='Helvetica-Bold', fontSize=11, leading=14, textColor=TEAL)

def marco(canv, doc):
    canv.saveState()
    canv.setFillColor(PALE); canv.rect(0, A4[1]-12*mm, A4[0], 12*mm, stroke=0, fill=1)
    canv.setFillColor(MUTED); canv.setFont('Helvetica', 7)
    canv.drawString(15*mm, A4[1]-8*mm, 'Documentación del aval · Curso virtual SEDA · uso interno')
    canv.drawRightString(A4[0]-15*mm, A4[1]-8*mm, 'Página %d' % doc.page)
    canv.setStrokeColor(MINT); canv.setLineWidth(.6); canv.line(15*mm, 12*mm, A4[0]-15*mm, 12*mm)
    canv.setFillColor(MUTED); canv.setFont('Helvetica', 6.5)
    canv.drawString(15*mm, 8.5*mm, 'Contiene datos personales de los docentes. No distribuir fuera de la coordinación.')
    canv.drawRightString(A4[0]-15*mm, 8.5*mm, 'Actualizado el 5 de octubre de 2026')
    canv.restoreState()

doc=BaseDocTemplate('SEDA_Directorio_Ponentes.pdf', pagesize=A4,
    leftMargin=15*mm, rightMargin=15*mm, topMargin=17*mm, bottomMargin=15*mm,
    title='Documentación del aval — directorio de ponentes', author='SEDA')
doc.addPageTemplates([PageTemplate(id='n',
    frames=[Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id='f')], onPage=marco)])
W=doc.width

w=openpyxl.load_workbook('SEDA_Matriz_Ponentes_Curso_Obesidad.xlsx')
m=w['Matriz 28 clases']; p=w['Ponentes - datos aval']
mods=['M1']*4+['M2']*4+['M3']*4+['M4']*4+['M5']*4+['M6']*4+['M7']*4

# --- emparejamiento de nombres -------------------------------------------
# No basta con el ultimo apellido: hay dos Molina, dos Cabrera y dos Gonzalez
# en la lista, y asi las clases se cruzaban de persona. Se comparan los
# nombres completos token a token, tolerando variantes de escritura
# (Janeth/Janneth, Palacios/Palacio) y se asigna cada clase a quien mas
# coincide.
TITULOS={'dr','dra','lcda','lcdo','msc','m','sc','mg','mgtr','de','del','la','y'}
def tokens(s):
    s=unicodedata.normalize('NFKD', str(s)).encode('ascii','ignore').decode().lower()
    return {t for t in re.findall(r'[a-z]+', s) if t not in TITULOS and len(t)>1}
def parecido(a,b):
    return a==b or SequenceMatcher(None,a,b).ratio()>=0.85
def puntaje(ta,tb):
    return sum(1 for a in ta if any(parecido(a,b) for b in tb))

clases={}
i=0
for r in range(3,31):
    if not m.cell(r,3).value: continue
    cod='%s·C%d'%(mods[i], i%4+1); i+=1
    if str(m.cell(r,6).value)=='VACANTE': continue
    clases.setdefault(str(m.cell(r,5).value), []).append(cod)

gente=[]
r=4
si=lambda v: str(v).strip().upper()=='SÍ'
while p.cell(r,2).value:
    n=str(p.cell(r,2).value)
    gente.append({'n':n,'esp':p.cell(r,3).value,'tel':p.cell(r,6).value,'mail':p.cell(r,5).value,
                  'ced':p.cell(r,7).value,'cv':si(p.cell(r,10).value),'fo':si(p.cell(r,11).value),
                  'co':si(p.cell(r,12).value),
                  'cl':''})
    r+=1

# cada clase se adjudica al docente cuyo nombre mas se le parece
for nombre_clase, cods in clases.items():
    tc=tokens(nombre_clase)
    mejor, mejor_p = None, 0
    for g in gente:
        p=puntaje(tc, tokens(g['n']))
        if p>mejor_p: mejor, mejor_p = g, p
    if mejor is None:
        raise SystemExit('Sin docente para la clase de %s' % nombre_clase)
    mejor.setdefault('cods', []).append((cods, nombre_clase))
for g in gente:
    cods=[c for par in g.get('cods', []) for c in par[0]]
    g['cl']=', '.join(sorted(cods)) if cods else '—'

tot_cv=sum(1 for g in gente if g['cv']); tot_fo=sum(1 for g in gente if g['fo'])
tot_co=sum(1 for g in gente if g['co']); tot_ml=sum(1 for g in gente if g['mail'])
N=len(gente)

F=[Paragraph('Documentación del aval', st_tit), Spacer(1,2.5*mm),
   Paragraph('Qué le falta a cada docente · Curso virtual SEDA de obesidad, diabetes, '
             'nutrición clínica y salud digital', st_sub), Spacer(1,4*mm)]

intro=Table([[Paragraph(
  '<b>Para qué sirve esta lista.</b> La Universidad de Cuenca exige, de cada uno de los %d docentes, '
  'cuatro documentos: hoja de vida, fotografía profesional, copia de la cédula o pasaporte y la '
  'declaración de conflicto de interés firmada. Sin eso no hay expediente que entregar.<br/><br/>'
  '<b>Cómo ayudar.</b> Lo más rápido es que cada persona lo suba al formulario; el enlace está abajo y '
  'se tarda cinco minutos. A quien no responda, se le escribe por WhatsApp: el teléfono está en la '
  'lista. <b>Las fotografías son lo más urgente</b> — si en el archivo de SEDA hay retratos de jornadas '
  'o simposios anteriores, sirven y resuelven de un golpe la columna peor parada.<br/><br/>'
  '<b>Todo se recibe también en</b> info@draomidresperez.com, en el formato que sea.' % N, st_av)]],
  colWidths=[W])
intro.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),WARM),('LINEBEFORE',(0,0),(0,-1),2.2,GOLD),
    ('LEFTPADDING',(0,0),(-1,-1),8),('RIGHTPADDING',(0,0),(-1,-1),8),
    ('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
F += [intro, Spacer(1,3*mm)]

link=Table([[Paragraph('Formulario para subir los documentos', st_hd)],
            [Paragraph('https://forms.gle/Xg7JRjNWqTNDeZWL8', st_big)]], colWidths=[W])
link.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),PALE),('ALIGN',(0,0),(-1,-1),'CENTER'),
    ('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4)]))
F += [link, Spacer(1,4*mm)]

res=Table([[Paragraph('<b>Cómo vamos</b>', st_av),
            Paragraph('Hojas de vida<br/><b>%d de %d</b>'%(tot_cv,N), st_av),
            Paragraph('Fotografías<br/><b>%d de %d</b>'%(tot_fo,N), st_av),
            Paragraph('Declaraciones<br/><b>%d de %d</b>'%(tot_co,N), st_av),
            Paragraph('Correos<br/><b>%d de %d</b>'%(tot_ml,N), st_av)]],
          colWidths=[W*0.2]+[W*0.2]*4)
res.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),VERDE),('ALIGN',(1,0),(-1,-1),'CENTER'),
    ('VALIGN',(0,0),(-1,-1),'MIDDLE'),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5),
    ('LEFTPADDING',(0,0),(-1,-1),6)]))
F += [res, Spacer(1,5*mm)]

hdr=Table([[Paragraph('DOCENTE', st_sec), Paragraph('CLASE', st_sec),
            Paragraph('TELÉFONO', st_sec), Paragraph('LE FALTA', st_sec)]],
          colWidths=[W*0.31, W*0.10, W*0.20, W*0.39])
hdr.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),TEAL),
    ('LEFTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4)]))
F.append(hdr)

filas=[]
for g in sorted(gente, key=lambda x: (x['cv']+x['fo']+x['co'], x['n'])):
    falta=[]
    if not g['cv']: falta.append('hoja de vida')
    if not g['fo']: falta.append('fotografía')
    if not g['ced']: falta.append('cédula/pasaporte')
    if not g['co']: falta.append('declaración firmada')
    if not g['mail']: falta.append('correo electrónico')
    izq=[Paragraph(g['n'], st_nom)]
    if g['esp']: izq.append(Paragraph(str(g['esp']), st_dat))
    cel=Paragraph(', '.join(falta), st_fal) if falta else Paragraph('nada — completo', st_ok)
    filas.append([izq, Paragraph(g['cl'], st_dat), Paragraph(str(g['tel'] or '—'), st_dat), cel])
t=Table(filas, colWidths=[W*0.31, W*0.10, W*0.20, W*0.39], repeatRows=0)
t.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),
    ('TOPPADDING',(0,0),(-1,-1),4.5),('BOTTOMPADDING',(0,0),(-1,-1),4.5),
    ('LEFTPADDING',(0,0),(-1,-1),5),('LINEBELOW',(0,0),(-1,-1),.35,MINT),
    ('ROWBACKGROUNDS',(0,0),(-1,-1),[colors.white, colors.HexColor('#FAFCFC')])]))
F.append(t)

F += [Spacer(1,5*mm), Paragraph(
  '<b>Dos personas que no hay que perseguir todavía:</b> la Dra. Gabriela Machado no ha confirmado su '
  'participación y la Dra. María Paz Castillo aún no ha sido invitada formalmente. De esas dos se encarga '
  'la coordinación.', st_av)]
doc.build(F)
print('generado')
