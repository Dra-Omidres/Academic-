# -*- coding: utf-8 -*-
"""BORRADOR de resumen curricular del Dr. Ramón de Cangas Morán (13/10/2026).

Lo pidió la Dra. Omidres para el expediente del aval, como candidato a M5 · C2. Él todavía
NO ha aceptado la clase: este documento no va a la Universidad hasta que acepte y apruebe el
texto.

Fuente: búsqueda web del 13/10/2026. Su página (ramondecangas.com) no se pudo abrir desde
aquí, así que todo sale de perfiles públicos: su fundación, Nutrigroup, Casa del Libro,
el registro de biografías de Asturias y EcuRed. Lo que las fuentes contradicen o que solo
aparece en una va POR CONFIRMAR. PubMed no devolvió artículos indexados a su nombre.
"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX

TEAL=RGBColor(0x0D,0x5A,0x62); GOLD=RGBColor(0xB8,0x85,0x2B); MUTED=RGBColor(0x5F,0x70,0x73)
d=Document()
for s in d.sections:
    s.top_margin=Cm(2); s.bottom_margin=Cm(2); s.left_margin=Cm(2.3); s.right_margin=Cm(2.3)
    h=s.header.paragraphs[0]; r=h.add_run('BORRADOR · pendiente de aceptación y aprobación del docente')
    r.bold=True; r.font.size=Pt(8); r.font.color.rgb=RGBColor(0xA4,0x16,0x1A)
st=d.styles['Normal']; st.font.name='Arial'; st.font.size=Pt(10.5)

def p(txt, b=False, color=None, size=None, despues=4):
    q=d.add_paragraph(); q.paragraph_format.space_after=Pt(despues)
    for i,trozo in enumerate(txt.split('[[')):
        if i:
            pc,resto=trozo.split(']]',1)
            r=q.add_run(pc); r.bold=True; r.font.highlight_color=WD_COLOR_INDEX.YELLOW; trozo=resto
        r=q.add_run(trozo); r.bold=b
        if color: r.font.color.rgb=color
        if size: r.font.size=Pt(size)
    return q
def sec(t): p(t, True, TEAL, 11.5, 3).paragraph_format.space_before=Pt(10)
def item(t):
    q=d.add_paragraph(style='List Bullet'); q.paragraph_format.space_after=Pt(2)
    for i,trozo in enumerate(t.split('[[')):
        if i:
            pc,resto=trozo.split(']]',1); r=q.add_run(pc); r.bold=True; r.font.highlight_color=WD_COLOR_INDEX.YELLOW; trozo=resto
        q.add_run(trozo)

p('RESUMEN CURRICULAR', True, GOLD, 9.5, 0)
p('Dr. Ramón de Cangas Morán', True, TEAL, 17, 2)
p('Dietista-nutricionista · Biólogo · Doctor en Biología Funcional y Molecular · Oviedo, España', False, MUTED, 10, 8)

sec('Formación académica')
item('Dietista-nutricionista — Universidad de Navarra, España. [[POR CONFIRMAR: nombre exacto del título y año.]]')
item('Licenciado en Biología — Universidad de Oviedo, España.')
item('Doctor en Biología Funcional y Molecular — Universidad de Oviedo. Tesis defendida el 13 de enero de 2011, sobre el efecto de un ingrediente de la dieta en la grasa corporal en personas con sobrepeso.')
item('Doctorado en Ciencia de los Alimentos — Universidad de La Habana, Cuba. [[POR CONFIRMAR: las fuentes no coinciden en si está concluido.]]')

sec('Cargos actuales')
item('Director de las clínicas de nutrición Nutrigroup (Oviedo, Gijón, Avilés, Madrid, Barcelona y Puebla, México).')
item('Presidente de la Fundación Alimenta Tu Salud, dedicada a la promoción de la salud a través de la nutrición.')

sec('Sociedades científicas')
item('Miembro de la Academia Española de Nutrición y Dietética.')
item('Miembro del comité científico de la Red Nacional Tu Nutricionista.')

sec('Líneas de trabajo')
item('Desarrollo de nuevos alimentos, microbiota intestinal y divulgación científica en nutrición.')

sec('Publicaciones de divulgación')
item('Microbiota y alimentación consciente. ISBN 978-84-415-4542-7.')
item('¡Come y ponte en forma! Desmontando los mitos de la nutrición deportiva.')
item('[[POR CONFIRMAR con él: artículos científicos indexados. PubMed no devolvió resultados a su nombre el 13/10/2026.]]')

sec('Docencia en este curso')
item('Curso virtual de actualización en obesidad, diabetes, nutrición clínica y salud digital — Sociedad de Endocrinología y Diabetes del Austro (SEDA), Ecuador. Módulo 5 · Clase 2: Estrategias nutricionales basadas en evidencia. [[Pendiente de su aceptación.]]')

p('')
p('Fuentes de este borrador: perfiles públicos consultados el 13/10/2026 (Fundación Alimenta Tu Salud, Nutrigroup, '
  'Casa del Libro, Reseñas biográficas de Asturias, EcuRed). Debe revisarlo y aprobarlo el docente antes de '
  'presentarlo a la Universidad.', False, MUTED, 8)
d.save('SEDA_Resumen_Curricular_Cangas_BORRADOR.docx')
print('generado')
